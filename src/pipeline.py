"""Load UCI Online Retail II into a reproducible DuckDB analytical model."""

from __future__ import annotations

import argparse
from collections.abc import Iterator
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB = ROOT / "data" / "retail.duckdb"
COLUMNS = {
    "Invoice": "invoice_no",
    "InvoiceNo": "invoice_no",
    "StockCode": "stock_code",
    "Description": "description",
    "Quantity": "quantity",
    "InvoiceDate": "invoice_at",
    "Price": "unit_price",
    "UnitPrice": "unit_price",
    "Customer ID": "customer_id",
    "CustomerID": "customer_id",
    "Country": "country",
}
REQUIRED = {"invoice_no", "stock_code", "quantity", "invoice_at", "unit_price", "country"}
INSERT_FIELDS = [
    "row_id", "source_sheet", "source_row", "invoice_no", "stock_code",
    "description", "quantity", "invoice_at", "unit_price", "customer_id", "country",
]


def _identifier(value: object) -> str | None:
    if pd.isna(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    result = str(value).strip()
    return result or None


def normalize(frame: pd.DataFrame, source_sheet: str, first_source_row: int,
              first_row_id: int) -> pd.DataFrame:
    """Map both UCI workbook schemas to a stable line-level schema."""
    renamed = frame.rename(columns=COLUMNS)
    if not REQUIRED.issubset(renamed.columns):
        missing = sorted(REQUIRED - set(renamed.columns))
        raise ValueError(f"Missing required source columns: {missing}")
    result = pd.DataFrame(index=range(len(renamed)))
    result["row_id"] = range(first_row_id, first_row_id + len(renamed))
    result["source_sheet"] = source_sheet
    result["source_row"] = range(first_source_row, first_source_row + len(renamed))
    for name in ("invoice_no", "stock_code", "customer_id"):
        result[name] = renamed[name].map(_identifier).values if name in renamed else None
    result["description"] = renamed.get("description", pd.Series([None] * len(renamed))).map(_identifier).values
    result["country"] = renamed["country"].map(_identifier).values
    result["quantity"] = pd.to_numeric(renamed["quantity"], errors="coerce").values
    result["unit_price"] = pd.to_numeric(renamed["unit_price"], errors="coerce").values
    result["invoice_at"] = pd.to_datetime(renamed["invoice_at"], errors="coerce").values
    return result[INSERT_FIELDS]


def iter_source(path: Path, chunk_size: int = 50_000) -> Iterator[tuple[str, int, pd.DataFrame]]:
    """Yield source rows without holding both workbook sheets in memory."""
    if path.suffix.lower() == ".csv":
        source_row = 2
        for frame in pd.read_csv(path, chunksize=chunk_size, low_memory=False):
            yield "CSV", source_row, frame
            source_row += len(frame)
    elif path.suffix.lower() == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows)
                batch = []
                source_row = 2
                for row in rows:
                    batch.append(row)
                    if len(batch) == chunk_size:
                        yield sheet.title, source_row, pd.DataFrame(batch, columns=headers)
                        source_row += len(batch)
                        batch = []
                if batch:
                    yield sheet.title, source_row, pd.DataFrame(batch, columns=headers)
        finally:
            book.close()
    else:
        raise ValueError("Input must be an .xlsx workbook or .csv file")


def create_raw_table(conn: duckdb.DuckDBPyConnection) -> None:
    conn.execute("""
        CREATE TABLE raw_lines (
            row_id BIGINT PRIMARY KEY, source_sheet VARCHAR, source_row BIGINT,
            invoice_no VARCHAR, stock_code VARCHAR, description VARCHAR,
            quantity DOUBLE, invoice_at TIMESTAMP, unit_price DOUBLE,
            customer_id VARCHAR, country VARCHAR
        )
    """)


def append_frame(conn: duckdb.DuckDBPyConnection, frame: pd.DataFrame,
                 source_sheet: str, source_row: int, first_row_id: int) -> int:
    normalized = normalize(frame, source_sheet, source_row, first_row_id)
    conn.register("incoming_batch", normalized)
    try:
        conn.execute(
            "INSERT INTO raw_lines SELECT " + ", ".join(INSERT_FIELDS) + " FROM incoming_batch"
        )
    finally:
        conn.unregister("incoming_batch")
    return len(normalized)


def create_model(conn: duckdb.DuckDBPyConnection) -> None:
    conn.execute((ROOT / "sql" / "model.sql").read_text(encoding="utf-8"))


def build_database(source: Path, destination: Path = DEFAULT_DB) -> int:
    """Create a fresh database; never replace an existing analytical result."""
    if not source.is_file():
        raise FileNotFoundError(source)
    if destination.exists():
        raise FileExistsError(f"Database already exists: {destination}. Choose a new --db path.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    conn = duckdb.connect(str(destination))
    total = 0
    try:
        conn.begin()
        create_raw_table(conn)
        for sheet, first_source_row, frame in iter_source(source):
            count = append_frame(conn, frame, sheet, first_source_row, total + 1)
            total += count
            print(f"Loaded {total:,} lines ({sheet})", flush=True)
        if total == 0:
            raise ValueError("Source contains no data rows")
        create_model(conn)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return total


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="UCI .xlsx or equivalent .csv")
    parser.add_argument("--db", type=Path, default=DEFAULT_DB, help="New DuckDB file")
    args = parser.parse_args()
    count = build_database(args.input, args.db)
    print(f"Created {args.db} with {count:,} source lines")


if __name__ == "__main__":
    main()
